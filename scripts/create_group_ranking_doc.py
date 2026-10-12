from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


OUT = "Group_Ranking_Worksheet.docx"


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(10)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)

styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"].font.size = Pt(10)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("Within-Group Creativity Ranking Worksheet")
run.bold = True
run.font.name = "Arial"
run.font.size = Pt(16)

intro = doc.add_paragraph()
intro.add_run(
    "Use this sheet after the annotator has independently scored each proposal. "
    "The creativity evaluation form scores are the primary data and should be calibrated against all proposals evaluated so far. "
    "This worksheet only records the relative order within each current group."
)

meta = doc.add_table(rows=2, cols=4)
meta.alignment = WD_TABLE_ALIGNMENT.CENTER
meta.style = "Table Grid"
labels = ["Annotator ID", "", "Date", ""]
for i, label in enumerate(labels):
    set_cell_text(meta.cell(0, i), label, bold=i in (0, 2))
labels = ["Session / Notes", "", "", ""]
for i, label in enumerate(labels):
    set_cell_text(meta.cell(1, i), label, bold=i == 0)
meta.cell(1, 1).merge(meta.cell(1, 3))

doc.add_paragraph()

for group_num in range(1, 10):
    if group_num in (4, 7):
        doc.add_page_break()

    heading = doc.add_paragraph()
    heading.style = styles["Heading 2"]
    heading_run = heading.add_run(f"Group {group_num}")
    heading_run.font.name = "Arial"
    heading_run.font.size = Pt(12)
    heading_run.bold = True

    table = doc.add_table(rows=6, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.columns[0].width = Inches(1.8)
    table.columns[1].width = Inches(5.6)

    rows = [
        ("Proposal IDs in this group", ""),
        ("Most creative proposal", ""),
        ("Middle proposal", ""),
        ("Least creative proposal", ""),
        ("Tie? If yes, explain", ""),
        ("Brief reason for current-group ranking", ""),
    ]
    for row_idx, (label, value) in enumerate(rows):
        set_cell_text(table.cell(row_idx, 0), label, bold=True)
        set_cell_text(table.cell(row_idx, 1), value)

    if group_num not in (3, 6, 9):
        doc.add_paragraph()

doc.save(OUT)
print(OUT)
