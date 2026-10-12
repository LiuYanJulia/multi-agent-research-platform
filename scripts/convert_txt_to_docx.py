from pathlib import Path

from docx import Document
from docx.shared import Pt


INPUT_DIR = Path("extracted_writings_by_user")


def convert_one(txt_path: Path) -> Path:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)

    text = txt_path.read_text(encoding="utf-8")
    lines = text.splitlines()

    if not lines:
        doc.add_paragraph("")
    else:
        for line in lines:
            doc.add_paragraph(line)

    out_path = txt_path.with_suffix(".docx")
    doc.save(out_path)
    return out_path


def main() -> None:
    txt_files = sorted(INPUT_DIR.glob("*.txt"))
    for txt_path in txt_files:
        out_path = convert_one(txt_path)
        print(out_path)
    print(f"Converted {len(txt_files)} files.")


if __name__ == "__main__":
    main()
