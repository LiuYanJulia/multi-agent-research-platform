from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("pilot_participant_introductions/Pilot_Study_Post_Task_Questionnaire.docx")

NAVY = "17365D"
PALE_BLUE = "EAF2F8"
LIGHT_GRAY = "D9D9D9"
TEXT = "1F2937"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=LIGHT_GRAY, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margin(cell, top=90, start=110, bottom=90, end=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:cantSplit")
    tr_pr.append(marker)


def set_font(run, size=10.5, bold=False, color=TEXT, italic=False):
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def clear_cell(cell):
    p = cell.paragraphs[0]
    for r in p.runs:
        p._p.remove(r._element)
    return p


def set_table_widths(table, widths):
    table.autofit = False
    grid_cols = table._tbl.tblGrid.gridCol_lst
    for grid_col, width in zip(grid_cols, widths):
        grid_col.set(qn("w:w"), str(int(width * 1440)))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    set_font(run, size=14 if level == 1 else 11.5, bold=True, color="000000")
    return p


def add_body(doc, text, bold_prefix=None, after=5):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix and text.startswith(bold_prefix):
        run = p.add_run(bold_prefix)
        set_font(run, bold=True)
        run = p.add_run(text[len(bold_prefix):])
        set_font(run)
    else:
        run = p.add_run(text)
        set_font(run)
    return p


def add_likert_table(doc, items):
    table = doc.add_table(rows=1, cols=8)
    table.style = "Table Grid"
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_table_widths(table, [3.57] + [0.43] * 7)

    headers = ["Statement", "1", "2", "3", "4", "5", "6", "7"]
    for cell, label in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
        set_cell_margin(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = clear_cell(cell)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if label == "Statement" else WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(label)
        set_font(run, size=8.5, bold=True, color="FFFFFF")
    set_repeat_table_header(table.rows[0])

    for index, item in enumerate(items):
        row = table.add_row()
        prevent_row_split(row)
        for cell in row.cells:
            set_cell_border(cell)
            set_cell_margin(cell, top=100, bottom=100)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if index % 2 == 1:
            for cell in row.cells:
                set_cell_shading(cell, "F8FAFC")
        p = clear_cell(row.cells[0])
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(item)
        set_font(run, size=9.2)
        for cell in row.cells[1:]:
            p = clear_cell(cell)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run("□")
            set_font(run, size=11)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_choice_line(doc, text, options):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text + " ")
    set_font(run, bold=True)
    for index, option in enumerate(options):
        run = p.add_run(f"□ {option}")
        set_font(run)
        if index < len(options) - 1:
            run = p.add_run("    ")
            set_font(run)
    return p


def add_response_lines(doc, count=3):
    for _ in range(count):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        run = p.add_run("________________________________________________________________________________")
        set_font(run, size=9.5, color="6B7280")


def add_question_with_lines(doc, number, question, lines=2):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(f"{number}. ")
    set_font(run, bold=True)
    run = p.add_run(question)
    set_font(run)
    add_response_lines(doc, lines)


def setup_doc():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.63)
    section.left_margin = Inches(0.68)
    section.right_margin = Inches(0.68)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(TEXT)

    for style_name in ("Title", "Heading 1", "Heading 2"):
        styles[style_name].font.name = "Aptos Display" if style_name == "Title" else "Aptos"
        styles[style_name]._element.rPr.rFonts.set(qn("w:ascii"), styles[style_name].font.name)
        styles[style_name]._element.rPr.rFonts.set(qn("w:hAnsi"), styles[style_name].font.name)
        styles[style_name].font.color.rgb = RGBColor(0, 0, 0)
    title_ppr = styles["Title"]._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Pilot Study Post Task Questionnaire")
    set_font(run, size=8, color="6B7280")
    return doc


def main():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = setup_doc()

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run("Pilot Study Post Task Questionnaire")
    set_font(run, size=22, bold=True, color="000000")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(16)
    run = p.add_run("Please complete this questionnaire after submitting your main task")
    set_font(run, size=10.5, color="4B5563")

    info = doc.add_table(rows=2, cols=2)
    info.style = "Table Grid"
    set_table_widths(info, [1.5, 5.0])
    fields = [("Participant ID", "____________________________________________"), ("Date", "____________________________________________")]
    for row, (label, blank) in zip(info.rows, fields):
        for cell in row.cells:
            set_cell_border(cell)
            set_cell_margin(cell, top=110, bottom=110)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(row.cells[0], PALE_BLUE)
        p = clear_cell(row.cells[0])
        run = p.add_run(label)
        set_font(run, size=9.5, bold=True)
        p = clear_cell(row.cells[1])
        run = p.add_run(blank)
        set_font(run, size=9.5, color="6B7280")

    add_heading(doc, "How to complete this questionnaire")
    add_body(doc, "Thank you for taking part in the pilot study. Your feedback will help us improve the platform and the study procedure. There are no right or wrong answers. Please answer based on your own experience.")
    add_body(doc, "For the rating tables, check one box in each row. The scale is 1 = Strongly disagree, 4 = Neither agree nor disagree, and 7 = Strongly agree.")
    add_body(doc, "Some questions refer to an additional AI guidance notice that may have appeared above the chat. If you did not see one, follow the skip instructions in that section.")

    add_heading(doc, "Part 1 Your experience with the AI chat")
    add_body(doc, "In this section, “AI chat responses” means the ordinary replies sent after you wrote a message. Do not include the separate additional AI guidance notice in your answers to these questions.")
    add_body(doc, "Please rate each statement.", after=4)
    core_items = [
        "The ordinary AI chat responses helped me make progress on the design task.",
        "The ordinary AI chat responses were relevant to what I was working on.",
        "It was easy to use the AI chat while working on the task.",
        "I could judge when an ordinary AI chat response was useful.",
        "I retained control over the direction of my final response.",
        "The ordinary AI chat responses supported a smooth flow of work.",
    ]
    add_likert_table(doc, core_items)

    add_heading(doc, "Part 2 Additional AI guidance")
    add_body(doc, "In this section, “additional AI guidance” means the separate notice that appeared above the chat with an Insert button. It does not refer to the ordinary AI chat responses.")
    add_body(doc, "Please answer this section only if you saw this notice during the activity. If you did not see one, check the relevant option below and continue to Part 3.")
    add_choice_line(doc, "Did you notice an additional AI guidance notice above the chat?", ["Yes", "No", "I do not remember"])
    add_choice_line(doc, "Approximately how many additional guidance notices did you notice?", ["0", "1", "2", "3", "4 or more", "I do not remember"])
    add_body(doc, "If you selected Yes above, please rate each statement below. Otherwise, continue to Part 3.", after=4)
    add_heading(doc, "How often the guidance was present", level=2)
    frequency_items = [
        "Additional AI guidance appeared frequently during the activity.",
        "Additional AI guidance appeared at multiple points as my work developed.",
        "Additional AI guidance repeatedly prompted me to reconsider what to do next.",
    ]
    add_likert_table(doc, frequency_items)

    add_heading(doc, "How strongly the guidance influenced my work", level=2)
    magnitude_items = [
        "Additional AI guidance encouraged a substantial change in my current approach.",
        "Additional AI guidance encouraged me to explore directions clearly different from my existing ideas.",
        "Additional AI guidance had a strong influence on the direction of my work.",
    ]
    add_likert_table(doc, magnitude_items)

    add_heading(doc, "Part 3 Task flow")
    add_body(doc, "If the writing area was unavailable at the beginning of your activity, please rate each statement below. Otherwise, continue to Part 4.", after=4)
    phase_items = [
        "The activity had a clear sequence: first explore ideas, then write them up.",
        "Having the writing area open later helped me focus on exploring ideas first.",
        "The delayed writing area prevented me from working in the way I wanted.",
    ]
    add_likert_table(doc, phase_items)

    add_heading(doc, "Part 4 Platform and study procedure")
    add_body(doc, "Please rate each statement. If you did not see additional guidance, leave the statements about the guidance notice and Insert action blank. If your writing area was available from the beginning, leave the final two statements blank.", after=4)
    pilot_items = [
        "Before starting, I understood what I needed to do.",
        "I understood what thinking aloud meant in this study.",
        "Speaking aloud was manageable while using the platform.",
        "It was clear whether my microphone was recording.",
        "It was clear how to continue a discussion with the AI.",
        "It was clear where and how to write my final response.",
        "The time available felt appropriate for the task.",
        "The additional guidance notice was visually easy to notice.",
        "The Insert action and what would happen next were clear.",
        "I understood why the writing area was initially unavailable.",
        "I knew when I could move to the writing stage.",
    ]
    add_likert_table(doc, pilot_items)

    add_heading(doc, "Part 5 Technical experience")
    add_choice_line(doc, "Was any part of the session technically disrupted?", ["No", "Yes"])
    add_body(doc, "If yes, check all that apply.", after=3)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run("□ Microphone     □ Transcription     □ Chat response     □ Writing area     □ Guidance notice     □ Timer     □ Submission     □ Other")
    set_font(run, size=9.5)
    add_body(doc, "If you selected Yes or Other, please describe what happened and when.", after=2)
    add_response_lines(doc, 3)
    add_choice_line(doc, "Did you refresh the page, navigate away, or lose any progress?", ["No", "Yes"])
    add_body(doc, "If yes, please describe what happened and what was lost.", after=2)
    add_response_lines(doc, 3)

    add_heading(doc, "Part 6 Your feedback")
    add_body(doc, "Please answer as fully as you can. Your comments will help us improve the platform before the main study.")
    open_questions = [
        ("What, if anything, was most helpful about working with the AI?", 3),
        ("What, if anything, was least helpful or most frustrating?", 3),
        ("Describe one moment when the AI changed, or did not change, what you did next.", 3),
        ("Please describe the moment you felt most unsure about what to do next.", 3),
        ("If you saw additional AI guidance, what did you think it was asking you to do? Was any wording confusing, too forceful, too vague, or poorly timed?", 4),
        ("Did thinking aloud change how you interacted with the AI or developed your ideas? In what way?", 3),
        ("If you could change one thing about the workspace or the AI interaction before another participant uses it, what would you change?", 4),
    ]
    for number, (question, lines) in enumerate(open_questions, 1):
        add_question_with_lines(doc, number, question, lines)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
