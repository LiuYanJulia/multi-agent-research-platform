from __future__ import annotations

import copy
import re
from pathlib import Path

from docx import Document


ROOT = Path("extracted_writings_by_user")
EVAL_DIR = ROOT / "creativity evaluation"
TEMPLATE = EVAL_DIR / "user01_creativity_evaluation.docx"


def user_ids_from_writings() -> list[str]:
    user_ids = set()
    for path in ROOT.glob("user*.docx"):
        match = re.match(r"(user\d+)_", path.name)
        if match:
            user_ids.add(match.group(1))
    return sorted(user_ids, key=lambda value: int(value.replace("user", "")))


def replace_first_paragraph_text(doc: Document, text: str) -> None:
    paragraph = doc.paragraphs[0]
    if paragraph.runs:
        first_run = paragraph.runs[0]
        first_run.text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def clone_core_properties(source: Document, target: Document) -> None:
    target.core_properties.author = source.core_properties.author
    target.core_properties.category = source.core_properties.category
    target.core_properties.comments = source.core_properties.comments
    target.core_properties.content_status = source.core_properties.content_status
    target.core_properties.identifier = source.core_properties.identifier
    target.core_properties.keywords = source.core_properties.keywords
    target.core_properties.language = source.core_properties.language
    target.core_properties.subject = source.core_properties.subject
    target.core_properties.version = source.core_properties.version


def main() -> None:
    if not TEMPLATE.exists():
        raise FileNotFoundError(TEMPLATE)

    template_doc = Document(TEMPLATE)
    users = user_ids_from_writings()

    for user_id in users:
        doc = Document(TEMPLATE)
        replace_first_paragraph_text(doc, f"{user_id} Creativity Evaluation")
        clone_core_properties(template_doc, doc)
        doc.core_properties.title = f"{user_id} Creativity Evaluation"
        out_path = EVAL_DIR / f"{user_id}_creativity_evaluation.docx"
        doc.save(out_path)
        print(out_path)

    print(f"Created or updated {len(users)} evaluation forms.")


if __name__ == "__main__":
    main()
